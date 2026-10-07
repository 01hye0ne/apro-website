# -*- coding: utf-8 -*-
"""국문영문글정리.xlsx '푸터' 시트 → apro-gray/policy-c.js (푸터 '개인정보처리방침' · '이메일무단수집거부' 팝업 글).

푸터 시트에서 자리표가 policy| 로 시작하는 줄이 팝업 글이다(policy|문서|순번|종류).
국문 = 변경 문구(있으면, '삭제'면 뺀다) 아니면 현재 문구, 영문 = 영문 문구(비면 영문판에서 뺀다).
종류 h 는 조 머리 — 그 앞의 줄은 머리말. 팝업 제목은 푸터 정책 링크 줄(자리표 footer|링크 이름 …)에서 가져온다.
영문 장은 아직 없어 팝업은 국문만 쓴다. 엑셀을 고친 뒤 다시 돌리면 된다:
  python content-export/tools/policy_js.py
"""
import io, json, sys
from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8")
BOOK = "content-export/국문영문글정리.xlsx"
OUT = "apro-gray/policy-c.js"
TITLE_EN = {"privacy": "Privacy Policy", "email": "Refusal of Unauthorized Email Collection"}


def txt(v):
    return str(v).strip() if v is not None and str(v).strip() else None


def main():
    ws = load_workbook(BOOK, read_only=True)["푸터"]
    docs = {k: {"ko": {"title": None, "intro": [], "arts": []}, "en": {"title": TITLE_EN[k], "intro": [], "arts": []}}
            for k in ("privacy", "email")}
    cur = {k: {"ko": None, "en": None} for k in docs}
    links = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        cur_t, new_t, en_t, loc = row[5], row[6], row[7], str(row[9] or "")
        if loc.startswith("footer|링크 이름"):
            links.append((txt(new_t) or txt(cur_t), txt(en_t)))
            continue
        if not loc.startswith("policy|"):
            continue
        _, doc, _i, kind = loc.split("|")
        ko = None if txt(new_t) == "삭제" else (txt(new_t) or txt(cur_t))
        for lang, t in (("ko", ko), ("en", txt(en_t))):
            if not t:
                continue
            d = docs[doc][lang]
            if kind == "h":
                cur[doc][lang] = {"t": t, "b": []}
                d["arts"].append(cur[doc][lang])
            elif cur[doc][lang] is None:
                d["intro"].append([kind, t])
            else:
                cur[doc][lang]["b"].append([kind, t])
    for k, (ko_title, en_title) in zip(("privacy", "email"), links):
        docs[k]["ko"]["title"] = ko_title
        if en_title:
            docs[k]["en"]["title"] = en_title
    js = ("/* 푸터 정책 팝업 글 — content-export/tools/policy_js.py 가 국문영문글정리.xlsx '푸터' 시트에서 만든다. 손으로 고치지 말 것.\n"
          "   [종류, 글] — h 조 머리 · s 굵은 머리글 · p 문단 · li 글머리 줄 */\n"
          "window.APRO_POLICY = " + json.dumps(docs, ensure_ascii=False, indent=1) + ";\n")
    io.open(OUT, "w", encoding="utf-8", newline="\n").write(js)
    for k, v in docs.items():
        print(k, "| ko '%s' 조 %d 머리말 %d" % (v["ko"]["title"], len(v["ko"]["arts"]), len(v["ko"]["intro"])),
              "| en '%s' 조 %d 머리말 %d" % (v["en"]["title"], len(v["en"]["arts"]), len(v["en"]["intro"])))


if __name__ == "__main__":
    main()

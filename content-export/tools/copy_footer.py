# -*- coding: utf-8 -*-
"""국문영문글정리.xlsx 에 '푸터' 시트를 더한다 — 푸터 글(메뉴 · 회사 정보 · 링크) + 정책 팝업 글 두 벌(국문 · 영문).

정책 글(개인정보처리방침 · 이메일무단수집거부)은 처음엔 _보관/04_푸터_정책문서.xlsx(옛 홈페이지 DB 에서 뽑은 것)에서 가져와
현재 문구 = 국문, 영문 문구 = 영문으로 한 줄에 짝지어 담는다. 다시 돌리면 앞 판 푸터 시트에서 사람이 채운 칸
(변경 · 영문 · 비고)을 현재 문구 기준으로 옮겨 담는다. 팝업 글은 이 시트에서 policy_js.py 가 만든다.
  python content-export/tools/copy_footer.py [앞 판 엑셀]
"""
import io, re, sys
from copy import copy
from lxml import html as LH
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding="utf-8")

BOOK = "content-export/국문영문글정리.xlsx"
POLICY_SRC = "content-export/_보관/04_푸터_정책문서.xlsx"
PAGE = "apro-gray/index-c.html"
SHEET = "푸터"
HEAD = ["순번", "섹션", "자리", "구분", "묶음", "현재 문구", "변경 문구", "영문 문구", "비고", "자리표(건드리지 마세요)"]
W = [6, 24, 16, 8, 6, 54, 54, 54, 20, 44]
HAN = re.compile(r"[가-힣ㄱ-ㆎ]")
ROLE = {"조제목": "조 제목", "소제목": "소제목", "본문": "본문", "목록": "목록"}
KEY = {"조제목": "h", "소제목": "s", "본문": "p", "목록": "li"}


def norm(s):
    return " ".join(str(s or "").split())


def footer_ui():
    """시안 C 푸터에 보이는 글 — 메뉴 · 회사 정보 · Copyright · 정책 링크 · 패밀리 사이트"""
    doc = LH.parse(io.StringIO(io.open(PAGE, encoding="utf-8").read())).getroot()
    f = doc.cssselect(".footer")[0]
    rows = []
    for col in f.cssselect(".f-nav .f-col"):
        head = norm(col.cssselect(".head")[0].text_content())
        rows.append(("푸터 메뉴 · " + head, "메뉴 제목", head, None, "GNB · 메가 메뉴 · 전체 메뉴의 대분류 이름과 같은 글"))
        for a in col.cssselect("a"):
            rows.append(("푸터 메뉴 · " + head, "메뉴 링크", norm(a.text_content()), None, "GNB · 메가 메뉴 · 전체 메뉴의 하위 이름과 같은 글"))
    for s in f.cssselect(".f-bottom .infos .line span"):
        rows.append(("푸터 회사 정보", "회사 정보", norm(s.text_content()), None, None))
    rows.append(("푸터 회사 정보", "저작권 표기", norm(f.cssselect(".f-legal .copy")[0].text_content()), None, None))
    for b in f.cssselect(".f-legal .f-pol"):
        rows.append(("푸터 정책 링크", "링크 이름 · 팝업 제목", norm(b.text_content()), b.get("data-policy"), None))
    rows.append(("푸터 패밀리 사이트", "단추", norm(f.cssselect(".fbtn span")[0].text_content()), None, None))
    for a in f.cssselect(".fmenu a"):
        rows.append(("푸터 패밀리 사이트", "목록", norm(a.text_content()), None, None))
    return rows


def policy_rows():
    """[섹션, 자리, 국문, 영문, 자리표, 비고] — 정책 두 문서를 한 줄씩 국문 · 영문 짝으로"""
    wb = load_workbook(POLICY_SRC, read_only=True)
    out = []
    ko = list(wb["개인정보처리방침_국문"].iter_rows(min_row=2, values_only=True))
    en = list(wb["개인정보처리방침_영문"].iter_rows(min_row=2, values_only=True))
    for i, (k, e) in enumerate(zip(ko, en)):      # 조 · 순서가 한 줄씩 맞는다(구분만 목록/본문으로 엇갈린 줄이 있음)
        art, ttl, kind, _n, text = k[:5]
        sec = "개인정보처리방침" + (" · %s %s" % (art, ttl) if art and art != "부칙" else " · 부칙" if art else " · 머리말")
        out.append([sec, ROLE.get(kind, "본문"), norm(text), norm(e[4]), "policy|privacy|%d|%s" % (i, KEY.get(kind, "p")), None])
    em = [r for r in wb["이메일무단수집거부"].iter_rows(min_row=2, values_only=True)]
    eko = [r for r in em if r[0] == "한국어"]
    een = [r for r in em if r[0] == "영문"]
    # 영문 첫 문단 = 국문 1 · 2 문단을 합친 것. 그 뒤는 한 줄씩 짝이 맞는다
    pair = {0: 0, 1: None}
    pair.update({i: i - 1 for i in range(2, len(eko))})
    for i, k in enumerate(eko):
        j = pair.get(i)
        note = "영문은 위 줄에 합쳐져 있음" if j is None else None
        out.append(["이메일무단수집거부", ROLE.get(k[1], "본문"), norm(k[3]), norm(een[j][3]) if j is not None else None,
                    "policy|email|%d|%s" % (i, KEY.get(k[1], "p")), note])
    return out


def carry(prev):
    keep = {}
    try:
        ws = load_workbook(prev)[SHEET]
    except Exception:
        return keep
    for row in ws.iter_rows(min_row=2, values_only=True):
        cur = norm(row[5])
        if cur:
            keep[cur] = {"new": row[6], "en": row[7], "note": row[8]}
    return keep


def build(prev=BOOK):
    old = carry(prev)
    wb = load_workbook(BOOK)
    if SHEET in wb.sheetnames:
        del wb[SHEET]
    pos = wb.sheetnames.index("홈") + 1 if "홈" in wb.sheetnames else 1
    ref = wb[wb.sheetnames[-1] if wb.sheetnames[-1] != "보관" else wb.sheetnames[-2]]
    ws = wb.create_sheet(SHEET, pos)

    thin = Side(style="thin", color="D0D5DB")
    bd = Border(left=thin, right=thin, top=thin, bottom=thin)
    yfill, bfill, gfill = (PatternFill("solid", fgColor=c) for c in ("FFF6D6", "E8F1FB", "F4F6F8"))
    ws.append(HEAD)
    for i, w in enumerate(W, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for c, rc in zip(ws[1], ref[1]):
        c.font, c.fill, c.alignment, c.border = copy(rc.font), copy(rc.fill), copy(rc.alignment), copy(rc.border)

    titles_en = {"privacy": "Privacy Policy", "email": "Refusal of Unauthorized Email Collection"}
    rows = []
    for sec, where, text, pol, note in footer_ui():
        en = titles_en.get(pol) if pol else (text if not HAN.search(text) else None)
        rows.append([sec, where, text, en, "footer|" + where + "|" + text, note])
    rows += policy_rows()

    n = 0
    for sec, where, text, en, loc, note in rows:
        n += 1
        got = old.get(norm(text), {})
        notes = [x for x in (note, got.get("note")) if x and x != note] if got.get("note") else ([note] if note else [])
        if note and note not in notes:
            notes.insert(0, note)
        ws.append([n, sec, where, "문구", None, text, got.get("new"), got.get("en") or en,
                   "\n".join(notes) or None, loc])
        row = ws[ws.max_row]
        for c in row:
            c.border = bd
            c.alignment = Alignment(vertical="top", wrap_text=True)
            c.font = Font(size=10)
        row[5].fill, row[6].fill, row[7].fill = gfill, yfill, bfill
        if where == "조 제목":
            row[5].font = Font(size=10, bold=True)
        row[9].font = Font(size=8, color="98A2B0")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:J%d" % ws.max_row

    g = wb["안내"]
    mark = "푸터 시트 (2026-10-07 추가)"
    if not any(c.value == mark for c in g["B"]):
        hs, bs = g["B3"], g["B4"]
        r = g.max_row + 2
        for txt, st in [(mark, hs),
                        ("푸터에 보이는 글(메뉴 · 회사 정보 · Copyright · 정책 링크 · 패밀리 사이트)과, 정책 링크를 누르면 뜨는 팝업 두 개"
                         "(개인정보처리방침 · 이메일무단수집거부)의 글 전부입니다. 정책 글은 옛 홈페이지에 있던 국문 · 영문을 한 줄씩 짝지어 넣었습니다.", bs),
                        ("정책 글을 고치면(노란 칸 · 파란 칸) 제가 팝업에 다시 넣습니다. 줄을 지우려면 변경 문구에 '삭제'라고만 써 주세요.", bs),
                        ("이메일무단수집거부의 영문 첫 문단은 국문 두 문단을 합친 글이라, 국문 둘째 줄의 영문 칸은 비워 두었습니다.", bs),
                        ("메뉴 이름은 GNB · 메가 메뉴 · 전체 메뉴와 같은 글입니다 — 한 곳을 바꾸면 모두 함께 고칩니다.", bs)]:
            c = g.cell(row=r, column=2, value=txt)
            c.font, c.alignment = copy(st.font), copy(st.alignment)
            r += 1
    wb.save(BOOK)
    print("푸터 시트 %d 줄 (푸터 글 %d · 정책 %d)" % (n, n - len(policy_rows()), len(policy_rows())))


if __name__ == "__main__":
    build(sys.argv[1] if len(sys.argv) > 1 else BOOK)

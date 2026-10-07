# -*- coding: utf-8 -*-
"""국문영문글정리.xlsx 에 나머지 페이지(시안 C) 시트를 더한다 — 사업영역 허브 · 회사정보 · 투자정보 · 지속가능경영 · 커뮤니티.

홈 · 푸터 · 사업영역 세부 네 장 · 게시글 상세(*-view-c)는 다른 생성기가 맡거나 빼는 장이다.
담지 않는 것: GNB · 메가 메뉴 · 전체 메뉴 · 푸터(푸터 시트에 있음), 게시판 글 목록 · 쪽 번호(게시판에서 관리),
글자 없이 숫자 · 기호뿐인 칸(재무 · 주식 표의 금액 등).

한 줄 = 화면의 글 한 덩이(제목 · 문단 · 목록 줄 · 표 칸 · 단추 …). 같은 장 안에서 같은 글은 한 줄로 합친다.
영문 칸은 옛 홈페이지에서 뽑아 둔 국문 · 영문 짝(content-export/data/*-texts.json, policies/*.md)에서
같은 국문이 있으면 미리 채우고, 영문뿐인 줄은 그대로 옮겨 둔다. 다시 돌리면 앞 판에서 사람이 채운 칸을 옮긴다.
  python content-export/tools/copy_pages.py [앞 판 엑셀]
"""
import glob, io, json, os, re, sys
from copy import copy
from lxml import html as LH
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding="utf-8")

BOOK = "content-export/국문영문글정리.xlsx"
PAGES = [  # (파일, 시트 이름)
    ("business-c", "사업영역 허브"),
    ("company-c", "회사 소개"), ("network-c", "글로벌 네트워크"), ("locations-c", "사업장소재"),
    ("ir-finance-c", "재무정보"), ("ir-disclosure-c", "공시정보"), ("ir-stock-c", "주식정보"), ("ir-policy-c", "내부규정"),
    ("esg-environment-c", "환경"), ("esg-social-c", "사회"), ("esg-governance-c", "지배구조"),
    ("esg-report-c", "내외부 신고채널"), ("esg-board-c", "ESG 게시판"),
    ("community-notice-c", "공지사항"), ("community-press-c", "보도자료"),
]
HEAD = ["순번", "섹션", "자리", "구분", "묶음", "현재 문구", "변경 문구", "영문 문구", "비고", "자리표(건드리지 마세요)"]
W = [6, 24, 16, 8, 6, 54, 54, 54, 20, 44]
HAN = re.compile(r"[가-힣ㄱ-ㆎ]")
LETTER = re.compile(r"[가-힣A-Za-z]")
SKIP_CLS = {"sf-gnb", "c-mega", "navover", "footer", "totop", "cpager", "sr-only", "skip"}
BLOCK = {"h1": "제목", "h2": "제목", "h3": "소제목", "h4": "항목 제목", "h5": "항목 제목", "h6": "항목 제목",
         "p": "문단", "li": "목록", "th": "표 머리", "td": "표 칸", "dt": "항목명", "dd": "항목값",
         "button": "단추", "label": "양식 항목", "summary": "펼침 단추", "caption": "표 제목", "option": "선택지",
         "figcaption": "그림 설명", "legend": "양식 묶음"}


def cls(e):
    return set((e.get("class") or "").split())


def norm(s):
    return " ".join(str(s or "").split())


def skipped(e):
    while e is not None:
        # 사이트 머리글(GNB)만 뺀다 — 콘텐츠 장 히어로도 <header class="cs-hero"> 라 header 전부를 빼면 안 된다
        if e.tag in ("footer", "script", "style", "svg", "noscript") or cls(e) & SKIP_CLS:
            return True
        # 게시판 글 목록 — 표 몸통 · 보도자료 목록 · 비어 있음 문구
        if e.tag == "tbody" and e.getparent() is not None and "cboard" in cls(e.getparent()):
            return True
        if "cnews" in cls(e) or "cempty" in cls(e):
            return True
        e = e.getparent()
    return False


def block_of(e):
    """글을 담는 가장 가까운 덩이 — 표 칸 · 목록 줄 · 문단 …, 없으면 그 원소. <a> 는 덩이 안이면 덩이를 따른다"""
    cur = e
    while cur is not None and cur.tag != "body":
        if cur.tag in BLOCK:
            return cur
        cur = cur.getparent()
    return e


def own_text(b, blocks):
    """덩이의 글 — 안에 다른 덩이가 있으면 그 글은 빼고, <br> 은 줄바꿈"""
    parts = [b.text or ""]
    def walk(el):
        for ch in el:
            if not isinstance(ch.tag, str):
                parts[-1] += ch.tail or ""
                continue
            if ch.tag == "br":
                parts.append(ch.tail or "")
                continue
            if ch in blocks or ch.tag in ("script", "style", "svg") or cls(ch) & {"sr-only"}:
                parts[-1] += ch.tail or ""
                continue
            # 붙어 있는 두 칸(<b>에이프로</b><span>262260</span> · 줄마다 나눈 <span>)은 화면에서 따로 보인다 — 줄을 나눈다.
            # 한 글자 칸(글자마다 나눈 등장 연출)과, 앞에 빈칸이 있는 강조(<b>전원 변환:</b> 고효율 …)는 그대로 잇는다
            inner = norm(ch.text_content())
            if len(inner) > 1 and parts[-1] and not parts[-1][-1].isspace():
                parts.append("")
            parts[-1] += ch.text or ""
            walk(ch)
            parts[-1] += ch.tail or ""
    walk(b)
    return "\n".join(" ".join(p.split()) for p in parts).strip()


def section_of(e):
    cur = e
    while cur is not None:
        c = cls(cur)
        if "cs-hero" in c or "ci-hero" in c or "ci-top" in c:
            return "히어로"
        if cur.tag == "section" or "cblock" in c or "cmodal" in c or "ci-band" in c:
            h = cur.cssselect("h2, h3")
            if h:
                t = norm(own_text(h[0], set()).replace("\n", " "))
                if "cmodal" in c:
                    return "팝업 · " + t
                return t[:40]
        cur = cur.getparent()
    return "본문"


def pull(path):
    tree = LH.parse(io.StringIO(io.open(path, encoding="utf-8", newline="").read()))
    body = tree.getroot().cssselect("body")[0]
    order, blocks = [], set()
    for e in body.iter():
        if not isinstance(e.tag, str) or skipped(e):
            continue
        if e.tag == "img":
            if norm(e.get("alt")):
                order.append(("alt", e))
            continue
        if norm(e.text) or any(norm(c.tail) for c in e if isinstance(c.tag, str)):
            b = block_of(e)
            if b not in blocks:
                blocks.add(b)
                order.append(("text", b))
    rows = []
    for kind, b in order:
        if kind == "alt":
            t, where, k = norm(b.get("alt")), "그림 대체글", "대체글"
        else:
            t = own_text(b, blocks)
            where = BLOCK.get(b.tag, "글")
            if b.tag == "a" or (b.tag not in BLOCK and b.getparent() is not None and b.getparent().tag == "a"):
                where = "링크"
            k = "문구"
        if not t or not LETTER.search(t):
            continue
        rows.append({"text": t, "where": where, "kind": k, "sec": section_of(b),
                     "loc": tree.getpath(b) + ("|alt" if kind == "alt" else "|text")})
    # 같은 장 안의 같은 글은 한 줄로
    merged, seen = [], {}
    for r in rows:
        key = (norm(r["text"]), r["kind"])
        if key in seen:
            m = seen[key]
            for f, v in (("secs", r["sec"]), ("wheres", r["where"])):
                if v not in m[f]:
                    m[f].append(v)
            m["locs"].append(r["loc"])
            continue
        m = {"text": r["text"], "kind": r["kind"], "secs": [r["sec"]], "wheres": [r["where"]], "locs": [r["loc"]]}
        seen[key] = m
        merged.append(m)
    return merged


# ── 옛 홈페이지 국문 · 영문 짝 ────────────────────────────────────────────
def en_map():
    out = {}
    def walk(k, e):
        if isinstance(k, str) and isinstance(e, str):
            for a, b in zip(k.split("\n"), e.split("\n")) if k.count("\n") == e.count("\n") else [(k, e)]:
                if HAN.search(a) and norm(b) and not HAN.search(b):
                    out.setdefault(norm(a), norm(b))
            if HAN.search(k) and norm(e) and not HAN.search(e):
                out.setdefault(norm(k), norm(e))
        elif isinstance(k, dict) and isinstance(e, dict):
            for key in k:
                if key in e:
                    walk(k[key], e[key])
        elif isinstance(k, list) and isinstance(e, list):
            for a, b in zip(k, e):
                walk(a, b)
    for f in glob.glob("content-export/data/*-texts.json"):
        d = json.load(io.open(f, encoding="utf-8"))
        if "KOR" in d and "ENG" in d:
            walk(d["KOR"], d["ENG"])
    # 정책 전문 md — 한국어 / English 두 덩이를 같은 차례의 글 줄끼리 짝짓는다
    pat = re.compile(r"^\s*(?:-\s*\*\*[^*]+\*\*:\s*|#{2,6}\s+)(.+)$")
    for f in glob.glob("content-export/policies/*.md"):
        s = io.open(f, encoding="utf-8").read()
        m = re.split(r"^##\s+(?:English|영문|ENG)\s*$", s, flags=re.M)
        if len(m) != 2:
            continue
        ko = [norm(x.group(1)) for x in map(pat.match, m[0].split("## 한국어", 1)[-1].splitlines()) if x]
        en = [norm(x.group(1)) for x in map(pat.match, m[1].splitlines()) if x]
        if len(ko) == len(en):
            for a, b in zip(ko, en):
                if HAN.search(a) and not HAN.search(b):
                    out.setdefault(a, b)
    return out


def carry(prev):
    keep = {}
    try:
        wb = load_workbook(prev)
    except Exception:
        return keep
    for _f, name in PAGES:
        if name not in wb.sheetnames:
            continue
        for row in wb[name].iter_rows(min_row=2, values_only=True):
            cur = norm(row[5])
            if cur:
                keep[(name, cur)] = {"new": row[6], "en": row[7], "note": row[8]}
    return keep


def build(prev=BOOK):
    old = carry(prev)
    enm = en_map()
    wb = load_workbook(BOOK)
    ref = wb["홈"]
    thin = Side(style="thin", color="D0D5DB")
    bd = Border(left=thin, right=thin, top=thin, bottom=thin)
    yfill, bfill, gfill = (PatternFill("solid", fgColor=c) for c in ("FFF6D6", "E8F1FB", "F4F6F8"))
    pos = wb.sheetnames.index("보관") if "보관" in wb.sheetnames else len(wb.sheetnames)
    total = pre = 0
    for f, name in PAGES:
        if name in wb.sheetnames:
            del wb[name]
            pos = wb.sheetnames.index("보관") if "보관" in wb.sheetnames else len(wb.sheetnames)
        ws = wb.create_sheet(name, pos)
        pos += 1
        ws.append(HEAD)
        for i, w in enumerate(W, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        for c, rc in zip(ws[1], ref[1]):
            c.font, c.fill, c.alignment, c.border = copy(rc.font), copy(rc.fill), copy(rc.alignment), copy(rc.border)
        rows = pull("apro-gray/%s.html" % f)
        for n, m in enumerate(rows, 1):
            t = m["text"]
            got = old.get((name, norm(t)), {})
            notes = []
            if len(m["locs"]) > 1:
                notes.append("같은 문구 %d 곳 — 한 줄로 합침" % len(m["locs"]))
            en = got.get("en")
            if not en:
                if norm(t) in enm:
                    en = enm[norm(t)]
                    pre += 1
                    notes.append("영문: 옛 홈페이지 글")
                elif not HAN.search(t):
                    en = t
            if got.get("note") and not str(got["note"]).startswith(("같은 문구", "영문: 옛")):
                notes.append(got["note"])
            ws.append([n, " · ".join(m["secs"]), " · ".join(m["wheres"]), m["kind"], None, t,
                       got.get("new"), en, "\n".join(notes) or None, "\n".join(m["locs"])])
            row = ws[ws.max_row]
            for c in row:
                c.border = bd
                c.alignment = Alignment(vertical="top", wrap_text=True)
                c.font = Font(size=10)
            row[5].fill, row[6].fill, row[7].fill = gfill, yfill, bfill
            if row[2].value and row[2].value.startswith(("제목", "소제목")):
                row[5].font = Font(size=10, bold=True)
            row[9].font = Font(size=8, color="98A2B0")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = "A1:J%d" % ws.max_row
        total += len(rows)
        print("  %-16s %4d 줄" % (name, len(rows)))

    g = wb["안내"]
    mark = "나머지 페이지 시트 (2026-10-07 추가)"
    if not any(c.value == mark for c in g["B"]):
        hs, bs = g["B3"], g["B4"]
        r = g.max_row + 2
        for txt, st in [(mark, hs),
                        ("사업영역 허브부터 보도자료까지 열다섯 장(시안 C)의 글입니다. 시트 하나가 한 장이고, 화면에 보이는 순서대로 담았습니다. "
                         "게시글 상세 장은 넣지 않았습니다.", bs),
                        ("넣지 않은 것: GNB · 메뉴 · 푸터(푸터 시트에 있음), 게시판 글 제목 · 날짜 · 쪽 번호(게시판에서 관리), "
                         "글자 없이 숫자뿐인 칸(재무 · 주식 표의 금액 등).", bs),
                        ("영문 칸 중 비고에 '영문: 옛 홈페이지 글'이라 적힌 줄은 옛 홈페이지에 있던 같은 국문의 영문을 미리 넣은 것입니다. "
                         "그대로 쓰셔도 되고 고치셔도 됩니다.", bs)]:
            c = g.cell(row=r, column=2, value=txt)
            c.font, c.alignment = copy(st.font), copy(st.alignment)
            r += 1
    wb.save(BOOK)
    print("열다섯 장 %d 줄 — 옛 홈페이지 영문 미리 채움 %d 줄" % (total, pre))


if __name__ == "__main__":
    build(sys.argv[1] if len(sys.argv) > 1 else BOOK)

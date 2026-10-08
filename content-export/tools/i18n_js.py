# -*- coding: utf-8 -*-
"""국문영문글정리.xlsx → apro-gray/i18n-en.js (시안 C 영문 보기용 국문 → 영문 사전).

KOR/ENG 를 누르면 같은 페이지에서 글만 바꾼다(시안 단계 방식, 2026-10-08). 페이지에 표식을 달지 않고
화면 글을 사전에서 찾아 바꾸므로, 엑셀의 '현재 문구'(그리고 '변경 문구')가 페이지 글과 같아야 걸린다.
여러 줄 문구(<br> 로 나뉜 것)는 줄 사이를 \n 으로 둔다. 같은 국문에 영문이 여럿이면 가장 많이 쓴 것.
정책 팝업 글(푸터 시트 policy|…)은 policy-c.js 가 따로 가진다. 엑셀을 고친 뒤 다시 돌린다:
  python content-export/tools/i18n_js.py
"""
import io, json, re, sys
from collections import Counter, defaultdict
from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8")
BOOK = "content-export/국문영문글정리.xlsx"
OUT = "apro-gray/i18n-en.js"
HAN = re.compile(r"[가-힣]")


def norm(v):
    """줄마다 공백을 하나로, 줄 사이는 \n"""
    lines = [" ".join(l.split()) for l in str(v or "").replace("\r", "").split("\n")]
    return "\n".join(l for l in lines if l)


def main():
    wb = load_workbook(BOOK, read_only=True)
    pick = defaultdict(Counter)
    for ws in wb.worksheets:
        if ws.title in ("안내", "보관"):
            continue
        for row in ws.iter_rows(min_row=2, values_only=True):
            if len(row) < 10:
                continue
            cur, new, en, loc = norm(row[5]), norm(row[6]), norm(row[7]), str(row[9] or "")
            if not en or loc.startswith("policy|"):
                continue
            for ko in (cur, new if new and new != "삭제" else ""):
                if not ko or not HAN.search(ko) or ko == en:
                    continue
                pick[ko][en] += 1
                # 여러 줄 문구는 한 줄로 이은 꼴과, 줄 수가 같으면 줄마다도 넣는다
                kl, el = ko.split("\n"), en.split("\n")
                if len(kl) > 1:
                    pick[" ".join(kl)][" ".join(el)] += 1
                    if len(kl) == len(el):
                        for a, b in zip(kl, el):
                            if HAN.search(a):
                                pick[a][b] += 1
    d = {ko: c.most_common(1)[0][0] for ko, c in pick.items()}
    js = ("/* 시안 C 영문 보기 사전(국문 → 영문) — content-export/tools/i18n_js.py 가 국문영문글정리.xlsx 에서 만든다. 손으로 고치지 말 것 */\n"
          "window.APRO_I18N_EN = " + json.dumps(d, ensure_ascii=False, indent=0, sort_keys=True) + ";\n")
    io.open(OUT, "w", encoding="utf-8", newline="\n").write(js)
    print("사전 %d 줄" % len(d))


if __name__ == "__main__":
    main()
